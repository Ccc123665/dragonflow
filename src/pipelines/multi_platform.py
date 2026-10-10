# -*- coding: utf-8 -*-
"""一题多发 · 选题转多平台内容。

用户的原话：「一个选题出多个平台内容」——这是内容工坊的第一条
「从选题开始」的流水线，前面几条都要求用户手上先有一篇现成的文章。

流程
  1. 联网检索：抓取搜索结果页的标题与摘要作为「事实底料」，保时效。
     检索失败自动降级为模型自身知识，并在日志里提示人工核对。
  2. 一次模型调用：基于检索结果 + 选题，产出多平台内容的结构化 JSON
     （小红书图文大纲 / 公众号长文 / 知乎头条观点文 / 短视频口播稿）。
  3. 按产出范围落地：
     · 图文 —— 复用 gzh_to_xhs.render_pages（kit_draw 排版出图，和
       「公众号链接转小红书图文」同一套版式引擎）
     · 文章 —— md + docx 双格式落盘（python-docx）
     · 视频 —— 复用 XhsVideo（分镜 + edge-tts 口播 + 逐帧渲染 + BGM）

为什么不做「先出拍板单再展开」：内容工坊的交互是「填参数 → 跑」，
参数本身就是确认。要改角度，改完参数再跑一次，比中途打断更省事。
"""
from __future__ import annotations

import os
import re
import urllib.parse
from typing import Any, Dict, List

from .base import Ctx, Input, Pipeline, register
from .gzh_to_xhs import render_pages
from .xhs_video import VOICES, XhsVideo

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/122.0 Safari/537.36")


# ================================================================ 联网检索
def _clean(s: str) -> str:
    import html as _html
    return _html.unescape(re.sub(r"<[^>]+>", "", s or "")).replace("\xa0", " ").strip()


def _sogou(query: str, limit: int) -> List[str]:
    """搜狗网页搜索。实测对免登录请求的结果相关度远好于必应（必应对
    无会话请求会把自然结果降级成查询首词的通用内容，cookie/RSS 都救不回来）。"""
    import requests

    r = requests.get("https://www.sogou.com/web?"
                     + urllib.parse.urlencode({"query": query}),
                     headers={"User-Agent": UA, "Referer": "https://www.sogou.com/",
                              "Accept-Language": "zh-CN,zh;q=0.9"}, timeout=15)
    r.encoding = "utf-8"
    page = r.text
    if "验证码" in page[:3000] or "antispider" in page[:3000]:
        return []
    out: List[str] = []
    # 每条结果从 <h3 class="vr-title"> 开始，取到下一个 vr-title 之前
    segments = re.split(r'<h3 class="vr-title[^"]*"', page)[1:]
    for seg in segments[:limit]:
        t = re.search(r"<a[^>]*>(.*?)</a>", seg, re.S)
        title = _clean(t.group(1)) if t else ""
        if not title:
            continue
        m = re.search(r'class="(?:text-layout|space-txt|star-wiki)[^"]*"[^>]*>(.*?)</(?:div|p)>',
                      seg, re.S)
        if not m:
            m = re.search(r"<p[^>]*>(.*?)</p>", seg, re.S)
        snip = re.sub(r"\s+", " ", _clean(m.group(1)))[:180] if m else ""
        out.append(f"· {title}" + (f"：{snip}" if snip else ""))
    return out


def _so360(query: str, limit: int) -> List[str]:
    """360 搜索，搜狗的备用引擎。"""
    import requests

    r = requests.get("https://www.so.com/s?" + urllib.parse.urlencode({"q": query}),
                     headers={"User-Agent": UA, "Referer": "https://www.so.com/",
                              "Accept-Language": "zh-CN,zh;q=0.9"}, timeout=15)
    r.encoding = "utf-8"
    out: List[str] = []
    for seg in re.split(r'<li class="res-list"', r.text)[1:limit + 1]:
        t = re.search(r"<h3[^>]*>\s*<a[^>]*>(.*?)</a>", seg, re.S)
        title = _clean(t.group(1)) if t else ""
        if not title:
            continue
        m = re.search(r'<p class="res-desc[^"]*"[^>]*>(.*?)</p>', seg, re.S)
        if not m:
            m = re.search(r"<p[^>]*>(.*?)</p>", seg, re.S)
        snip = re.sub(r"\s+", " ", _clean(m.group(1)))[:180] if m else ""
        out.append(f"· {title}" + (f"：{snip}" if snip else ""))
    return out


def web_search(query: str, ctx: Ctx, *, limit: int = 8) -> str:
    """多引擎检索「事实底料」：搜狗 → 360，全部失败返回空串（调用方降级）。

    为什么不用必应：实测必应对免登录的程序化请求会把自然结果替换成
    查询首词的通用内容（「广东选调生 机考」只回「广东省百科」），
    加 cookie、换 RSS 格式都一样，无法用于保时效。百度则直接出验证页。
    """
    for name, fn in (("搜狗", _sogou), ("360", _so360)):
        if ctx.stopped():
            return ""
        try:
            items = fn(query, limit)
        except Exception as e:
            ctx.log(f"{name}检索异常（{str(e)[:60]}）")
            items = []
        if items:
            ctx.log(f"{name}检索到 {len(items)} 条结果")
            return "\n".join(items)
        ctx.log(f"{name}没有可用结果")
    return ""


# ================================================================ 多平台生成
SYSTEM = (
    "你是一位全媒体内容主编，手下管着小红书、公众号、知乎、头条和短视频五条内容线。"
    "你写任何平台的内容都绝不使用「包过」「保过」「押题」「最」「第一」「100%」"
    "等违规承诺词与广告法极限词，不承诺任何结果，不写任何引流信息（微信号/二维码/链接）。"
    "你只输出 JSON，不要任何解释文字。"
)

PROMPT = """围绕选题「{topic}」，为多个平台各产出一版内容。

【检索到的事实底料】（{retrieved_note}）
{facts}

【账号人设】{persona}
【目标人群】{audience}
【补充要求】{extra}

事实纪律（最重要）：
1. 硬事实（数字、日期、政策、人名职务）只能来自上面的检索底料，一个都不许编造。
2. 检索底料没有覆盖到的内容，用常识性表述带过，不写成确定事实。
3. 易变信息（截止时间、名额、价格、政策状态）必须带时间限定，例如「截至{year}年」「以最新公告为准」。
4. 当前是{year}年。任何年份表述都要和这个对齐。

请输出一个 JSON 对象，不要加注释，不要 markdown 代码块以外的文字：

{{
  "xhs": {{
    "title": "小红书标题，不超过 20 字，带钩子但不夸张",
    "body": "正文 500-900 字，短句分段，3-4 个 emoji 小标题，口语化，结尾一句互动引导",
    "topics": "6-10 个话题标签，空格分隔，不带 # 号",
    "cover": "封面文案，不超过 12 字",
    "pages": [ ...共 {pages} 页... ]
  }},
  "gzh": {{
    "title": "公众号标题，不超过 30 字",
    "digest": "摘要，不超过 60 字，用于分享卡片",
    "body_md": "正文 1500-2500 字，Markdown：开头引入 + 3-4 个二级标题分段（每段先结论再原因）+ 结尾落点"
  }},
  "zhihu": {{
    "title": "知乎/头条标题，观点或悬念式，20-30 字",
    "body_md": "正文 1000-2000 字，Markdown：首段亮观点，2-3 层论证，回应一个常见反方观点，结尾收束"
  }},
  "video": {{
    "title": "短视频标题，不超过 30 字",
    "body": "口播稿 220-300 字：前 3 秒硬钩子，中段 3-4 个信息点短句快节奏，结尾收束不带引导话术。分段落写，每段就是一个分镜"
  }}
}}

pages 里每一项从下面 7 种版式里选一种，严格按字段填（这是排版引擎的规范，超字会被裁掉）：

1) 封面（必须第 1 页，type=cover）
   {{"type":"cover","tag":"左上角小标签≤8字","title_small":"小标题≤10字",
     "title_big":["大标题第一行≤7字","大标题第二行≤7字"],
     "accent":"第二行里要染色的关键词，2-4字，没有传空字符串",
     "sub1":"副标题≤18字","sub2":"补充说明≤22字",
     "stats":[["数字","单位","说明≤6字"]}}
   stats 最多 3 组，必须是检索底料里真实出现的数字，没有就 []。
2) 列表页 {{"type":"rows","tag":"≤8字","title":"≤12字","sub":"≤20字","banner":"≤28字",
   "rows":[["名称≤16字","数值胶囊≤6字","说明≤22字",false]]}} rows 3-6 条，true=重点行。
3) 表格页 {{"type":"table","tag":"≤8字","title":"≤12字","sub":"≤20字","banner":"≤28字",
   "headers":["列名"],"rows":[["第一列","第二列","…","none"]]}} 末项 none|hot|total。
4) 要点页 {{"type":"bullets","tag":"≤8字","title":"≤12字","sub":"≤20字","banner":"≤28字",
   "items":[["要点标题≤14字","补充说明≤30字"]]}} items 3-6 条。
5) 金句页 {{"type":"quote","tag":"≤8字","title":"≤12字","quote":"≤60字金句","source":"出处≤10字"}}
6) 段落页 {{"type":"text","tag":"≤8字","title":"≤12字","sub":"≤20字",
   "paras":["段落≤90字"]}} paras 最多 4 段。
7) 收尾页（必须最后一页，type=end）
   {{"type":"end","title":"≤10字","lines":["要点≤16字"],"cta":"行动号召≤16字"}}

硬性要求：pages 正好 {pages} 页，第 1 页 cover、最后一页 end，中间至少 2 页 rows/table；
所有数字必须来自检索底料；各平台之间事实一致、角度差异化（小红书口语化、公众号深度解读、
知乎观点鲜明、短视频强钩子）。
"""

RETRIEVED_OK = "来自实时检索，可直接引用，注意其中出现的日期"
RETRIEVED_NONE = ("检索不可用，以下为空。请只用常识性表述，所有具体数字/日期一律不写，"
                  "并在正文里提示读者以官方最新信息为准")


class _ScaledCtx:
    """把子步骤的进度重标到外层的一个区间，避免进度条来回倒退。

    复用的子流水线（render_pages / XhsVideo）内部都会按 0~1 报自己的进度，
    直接透传的话外层进度会从 78% 跳回 5%。包一层，把它们的 0~1 线性映射到
    指定区间；没显式覆写的方法一律转发给原 ctx（settings / out_dir / llm…）。
    """

    def __init__(self, ctx: Ctx, lo: float, hi: float):
        self._ctx = ctx
        self._lo, self._hi = float(lo), float(hi)
        self.artifacts = ctx.artifacts        # 共享同一份产物登记表

    def log(self, msg: str) -> None:
        self._ctx.log(msg)

    def progress(self, pct: float, msg: str = "") -> None:
        p = max(0.0, min(1.0, float(pct or 0)))
        self._ctx.progress(self._lo + (self._hi - self._lo) * p, msg)

    def stopped(self) -> bool:
        return self._ctx.stopped()

    def check(self) -> None:
        self._ctx.check()

    def add(self, path: str) -> str:
        return self._ctx.add(path)

    def llm(self, *args, **kw):
        return self._ctx.llm(*args, **kw)

    def __getattr__(self, name):
        return getattr(self._ctx, name)


# ================================================================ 文章落盘
def _md_to_docx(md_text: str, path: str) -> None:
    """极简 Markdown → docx：标题 / 列表 / 引用 / **粗体**。纯排版诉求，够用。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    ACCENT = RGBColor(0x3B, 0x6F, 0xE0)

    def set_font(run, size=None, bold=None, color=None):
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    st.font.size = Pt(12)

    for line in (md_text or "").splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            h = doc.add_heading("", level=min(len(m.group(1)), 4))
            set_font(h.add_run(m.group(2).replace("**", "")),
                     size={1: 20, 2: 16, 3: 14, 4: 13}[len(m.group(1))],
                     bold=True, color=ACCENT if len(m.group(1)) <= 2 else None)
            continue
        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            self_rich(p, m.group(1), set_font)
            continue
        m = re.match(r"^>\s*(.*)", s)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            self_rich(p, m.group(1), set_font, grey=True)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.4
        self_rich(p, s, set_font)

    doc.save(path)


def self_rich(par, text, set_font, size=12, grey=False):
    """解析 **粗体** 内联。"""
    from docx.shared import RGBColor
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size=size, bold=True)
        else:
            set_font(par.add_run(part), size=size,
                     color=RGBColor(0x66, 0x66, 0x66) if grey else None)


# ================================================================ Pipeline
@register
class MultiPlatform(Pipeline):
    key = "multi_platform"
    name = "一题多发 · 选题转多平台内容"
    desc = ("输入一个选题（而不是一篇现成文章），联网检索保时效，一次产出"
            "小红书图文（文案+配图）、公众号/知乎头条文章（md+docx）、9:16 竖屏短视频")
    note = ("全部产出约 3-6 分钟（检索 + 模型 + 出图 + 视频渲染约为成片时长的 4 倍）。"
            "检索走搜狗，备用 360，都失败会自动降级为模型知识并在日志提示——"
            "涉及数字和政策的选题，发布前请人工核对时效。中途可点「停止」。")
    platform = "xhs"
    content_type = "图文 + 文章 + 视频"
    icon = "sparkle"
    fields = ["title", "body", "topics", "cover", "image_hint"]
    outputs = ["小红书文案 + 配图", "公众号/知乎文章（md+docx）", "9:16 竖屏短视频（mp4）"]

    def inputs(self) -> List[Input]:
        return [
            Input("topic", "选题（可把详细要求一起写上）", "textarea", required=True,
                  rows=3, placeholder="例如：广东选调机考的科普，面向 2027 届应届生"),
            Input("scope", "产出范围", "choice", default="全部（图文+文章+视频）",
                  choices=["全部（图文+文章+视频）", "仅图文（小红书）",
                           "仅文章（公众号+知乎头条）", "仅视频（竖屏短视频）"],
                  help="只想改其中一种时，选单项能省一大半时间（视频渲染占大头）"),
            Input("persona", "账号人设", "text", required=False,
                  placeholder="留空则自动贴合选题"),
            Input("audience", "目标人群", "text", required=False,
                  placeholder="例如：在职备考的上班族"),
            Input("pages", "图文页数", "choice", default="8",
                  choices=["6", "8", "10"], required=False, help="含封面和收尾页"),
            Input("voice", "口播音色", "choice", default="女声 · 亲和",
                  choices=list(VOICES.keys()), required=False),
            Input("bgm", "背景音乐", "choice", default="轻和声（无版权）",
                  choices=["轻和声（无版权）", "不要"], required=False),
            Input("extra", "补充要求", "textarea", required=False, rows=2,
                  placeholder="例如：语气轻松一点，多讲待遇数字"),
        ]

    def run(self, values: Dict[str, str], ctx: Ctx) -> Dict[str, Any]:
        import llm

        topic = (values.get("topic") or "").strip()
        if len(topic) < 4:
            raise ValueError("先把选题写清楚（至少 4 个字）")
        scope = values.get("scope") or "全部（图文+文章+视频）"
        want_xhs = scope in ("全部（图文+文章+视频）", "仅图文（小红书）")
        want_article = scope in ("全部（图文+文章+视频）", "仅文章（公众号+知乎头条）")
        want_video = scope in ("全部（图文+文章+视频）", "仅视频（竖屏短视频）")
        pages_n = int(values.get("pages") or 8)
        persona = (values.get("persona") or "").strip() or "自动贴合选题"
        audience = (values.get("audience") or "").strip() or "自动贴合选题"
        extra = (values.get("extra") or "").strip() or "无"

        # ---- 1. 联网检索（搜狗 → 360 → 降级）----
        ctx.progress(0.04, "联网检索事实底料")
        import datetime
        year = datetime.date.today().year
        facts = web_search(f"{topic} {year}", ctx)
        if facts:
            # 第二轮：不加年份再补一批泛结果
            more = web_search(topic, ctx, limit=5)
            if more:
                facts += "\n" + more
        retrieved_note = RETRIEVED_OK if facts else RETRIEVED_NONE
        if not facts:
            ctx.log("⚠ 本次没有检索到实时资料，产出将基于模型常识——"
                    "数字和政策类信息请人工核实后再发布")

        # ---- 2. 一次模型调用，产多平台内容 ----
        ctx.progress(0.18, "生成多平台内容（一次调用）")
        raw = ctx.llm(SYSTEM, PROMPT.format(
            topic=topic, facts=facts or "（无）", retrieved_note=retrieved_note,
            persona=persona, audience=audience, extra=extra, pages=pages_n, year=year),
            max_tokens=12000, temperature=0.7)
        data = _json_of(raw)
        if not data:
            raise ValueError("模型没有返回可解析的 JSON，请重试或换一个模型")

        xhs = data.get("xhs") or {}
        gzh = data.get("gzh") or {}
        zhihu = data.get("zhihu") or {}
        video = data.get("video") or {}

        out: Dict[str, Any] = {
            "title": _clip(xhs.get("title"), 60) or topic[:60],
            "body": str(xhs.get("body") or "").strip(),
            "topics": _clip(xhs.get("topics"), 200),
            "cover": _clip(xhs.get("cover"), 30),
        }
        made: List[str] = []

        # ---- 3. 图文（复用 gzh_to_xhs 的排版渲染，换暖橙奶油杂志风画布）----
        if want_xhs:
            from .gzh_to_xhs import _normalize_pages
            from .gzh_to_xhs_magazine import MagazineDeck
            pages = _normalize_pages(xhs.get("pages"), pages_n)
            if pages:
                ctx.progress(0.32, f"出图（共 {len(pages)} 页）")
                out_dir = self.output_dir(topic)
                # 子步骤进度重标到 0.34~0.72，进度条不倒退
                made += render_pages(pages, out_dir, _ScaledCtx(ctx, 0.34, 0.72),
                                     deck_cls=MagazineDeck)
                ctx.log(f"图文配图已输出到 {out_dir}")
                out["image_hint"] = (f"已出图 {len(made)} 张，目录：{out_dir}")
            else:
                ctx.log("模型没有返回可用的图文页面大纲，图文跳过")
        if "image_hint" not in out:
            out["image_hint"] = ""

        # ---- 4. 文章（md + docx）----
        if want_article:
            ctx.progress(0.75, "文章落盘（md + docx）")
            art_dir = self.output_dir(topic)
            os.makedirs(art_dir, exist_ok=True)
            safe = _safe_name(out["title"] or topic)
            for tag, piece in (("公众号", gzh), ("知乎头条", zhihu)):
                md = _article_md(piece, tag)
                if not md:
                    ctx.log(f"{tag}正文为空，跳过")
                    continue
                p_md = os.path.join(art_dir, f"{safe}_{tag}.md")
                p_dx = os.path.join(art_dir, f"{safe}_{tag}.docx")
                with open(p_md, "w", encoding="utf-8") as f:
                    f.write(md)
                try:
                    _md_to_docx(md, p_dx)
                    ctx.log(f"{tag}：md + docx 已落盘")
                except Exception as e:
                    ctx.log(f"{tag} docx 生成失败（{str(e)[:60]}），md 已保留")
                ctx.add(p_md)
                made.append(p_md)
                if os.path.exists(p_dx):
                    ctx.add(p_dx)
                    made.append(p_dx)
            out["image_hint"] = (out.get("image_hint", "")
                                 + f"\n文章输出目录：{art_dir}").strip()

        # ---- 5. 短视频（复用 xhs_video 全套：分镜/TTS/渲染/BGM）----
        if want_video:
            v_body = str(video.get("body") or "").strip()
            if len(v_body) >= 20:
                ctx.progress(0.78, "渲染竖屏短视频")
                v = dict(values)
                v["title"] = _clip(video.get("title"), 60) or out["title"]
                v["body"] = v_body
                v["style"] = "cream"          # 视频画面与图文同一套奶油杂志风
                # 视频渲染占大头，进度重标到 0.80~0.98，收尾再置 1.0
                vres = XhsVideo().run(v, _ScaledCtx(ctx, 0.80, 0.98)) or {}
                v_shot = vres.get("shot_hint") or ""
                if v_shot:
                    out["image_hint"] = (out.get("image_hint", "")
                                         + f"\n分镜脚本：{str(v_shot)[:120]}").strip()
            else:
                ctx.log("模型没有返回可用的口播稿，视频跳过")

        ctx.progress(1.0, "完成")
        out["_artifacts"] = list(ctx.artifacts)
        return out


def _article_md(piece: Dict[str, Any], tag: str) -> str:
    """把 gzh/zhihu 两块 JSON 拼成一篇完整 md。"""
    title = _clip(piece.get("title"), 80)
    body = str(piece.get("body_md") or "").strip()
    if not body:
        return ""
    head = f"# {title or tag}"
    digest = _clip(piece.get("digest"), 120)
    if digest and tag == "公众号":
        head += f"\n\n> 摘要：{digest}"
    return head + "\n\n" + body + "\n"


def _safe_name(text: str) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff]+", "_", str(text or "")).strip("_")[:24] or "output"


def _clip(v: Any, n: int) -> str:
    return str(v if v is not None else "").strip()[:n]


def _json_of(raw: str) -> Dict[str, Any] | None:
    import json

    import llm
    obj = llm.extract_json(raw)
    if obj:
        return obj
    m = re.search(r"\{.*\}", raw or "", re.S)
    if m:
        try:
            return json.loads(m.group(0))
        except Exception:
            return None
    return None
